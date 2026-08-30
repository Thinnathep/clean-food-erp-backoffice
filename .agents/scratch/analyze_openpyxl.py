import openpyxl
import json

file_path = r"d:\VS CODE\ร้าน\Google Sheet\ระบบบัญชีร้านส้มตำ น้ำปั่น V.3\ระบบบัญชีร้านส้มตำซูวี_น้ำปั่น_V3.xlsx"

try:
    wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
    result = {}
    for sheet_name in wb.sheetnames:
        sheet = wb[sheet_name]
        # Get first row (headers)
        headers = []
        for row in sheet.iter_rows(min_row=1, max_row=1, values_only=True):
            headers = [str(h) if h is not None else "" for h in row]
            break
        result[sheet_name] = headers
    
    print(json.dumps(result, ensure_ascii=False, indent=2))
except Exception as e:
    print(f"Error: {e}")
